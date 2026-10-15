"""Excel input and output of the engine.

Input: a workbook with a long-format sheet `Precios` (VARIEDAD, INFORMANTE, TIPO, MES, PRECIO) and,
optionally, a sheet `Variedades` (VARIEDAD, plus VARIEDAD_12 or GRUPO_1..GRUPO_n, plus
TRATAMIENTO_ESPECIAL) that describes the imputation hierarchy.
"""
from __future__ import annotations

import io
import re
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook, load_workbook

from .domain.models import Observacion, Resultado, Variedad

SHEET_PRECIOS = "Precios"
SHEET_VARIEDADES = "Variedades"
_REQUIRED = ("VARIEDAD", "INFORMANTE", "TIPO", "MES", "PRECIO")
_TRUE = {"SI", "SÍ", "S", "1", "TRUE", "X", "VERDADERO"}
_MES = re.compile(r"^\d{4}-\d{2}")


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _mes(value: Any, row: int) -> str:
    if isinstance(value, (datetime, date)):
        return f"{value.year:04d}-{value.month:02d}"
    text = _text(value)
    if _MES.match(text):
        return text[:7]
    raise ValueError(f"fila {row}: el mes {text!r} no tiene el formato AAAA-MM")


def _precio(value: Any, row: int) -> float:
    if isinstance(value, bool):
        raise ValueError(f"fila {row}: el precio {value!r} no es un número")
    if isinstance(value, (int, float)):
        return float(value)
    text = _text(value).replace(",", ".")
    try:
        return float(text)
    except ValueError:
        raise ValueError(f"fila {row}: el precio {value!r} no es un número") from None


def _sheet(workbook: Any, name: str) -> Any | None:
    for candidate in workbook.sheetnames:
        if candidate.strip().lower() == name.lower():
            return workbook[candidate]
    return None


def _chain_from_code12(code: str) -> tuple[str, ...]:
    # The 12-division code is hierarchical: shorter prefixes are broader groups.
    return tuple(code[:n] for n in (len(code), 6, 5, 4, 3, 1) if 0 < n <= len(code))


def _read_variedades(sheet: Any) -> list[Variedad]:
    rows = sheet.iter_rows(values_only=True)
    header = [_text(h).upper() for h in next(rows, ())]
    if "VARIEDAD" not in header:
        raise ValueError(f"la hoja {SHEET_VARIEDADES} necesita la columna VARIEDAD")

    index = {name: i for i, name in enumerate(header) if name}
    grupo_cols = sorted(
        (name for name in index if re.fullmatch(r"GRUPO_\d+", name)),
        key=lambda name: int(name.split("_")[1]),
    )

    variedades: list[Variedad] = []
    for row in rows:
        if not row or all(c is None for c in row):
            continue
        get = lambda col: row[index[col]] if col in index and index[col] < len(row) else None  # noqa: E731
        codigo = _text(get("VARIEDAD"))
        if not codigo:
            continue

        grupos: tuple[str, ...] = ()
        if grupo_cols:
            grupos = tuple(g for g in (_text(get(c)) for c in grupo_cols) if g)
        elif _text(get("VARIEDAD_12")):
            grupos = _chain_from_code12(_text(get("VARIEDAD_12")))

        especial = _text(get("TRATAMIENTO_ESPECIAL")).upper() in _TRUE
        variedades.append(Variedad(codigo, grupos, especial))

    return variedades


def leer_entrada(contenido: bytes) -> tuple[list[Observacion], list[Variedad]]:
    try:
        workbook = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception:
        raise ValueError("el archivo no es un Excel (.xlsx) válido") from None

    sheet = _sheet(workbook, SHEET_PRECIOS)
    if sheet is None:
        raise ValueError(f"el Excel debe tener una hoja llamada {SHEET_PRECIOS}")

    rows = sheet.iter_rows(values_only=True)
    header = [_text(h).upper() for h in next(rows, ())]
    missing = [name for name in _REQUIRED if name not in header]
    if missing:
        raise ValueError(f"la hoja {SHEET_PRECIOS} no tiene las columnas: {', '.join(missing)}")
    index = {name: header.index(name) for name in _REQUIRED}

    observaciones: list[Observacion] = []
    for number, row in enumerate(rows, start=2):
        if not row or all(c is None for c in row):
            continue
        cell = lambda name: row[index[name]] if index[name] < len(row) else None  # noqa: E731
        variedad, informante = _text(cell("VARIEDAD")), _text(cell("INFORMANTE"))
        if not variedad or not informante:
            raise ValueError(f"fila {number}: faltan la variedad o el informante")
        observaciones.append(
            Observacion(
                variedad=variedad,
                informante=informante,
                tipo=_text(cell("TIPO")).upper(),
                mes=_mes(cell("MES"), number),
                precio=_precio(cell("PRECIO"), number),
            )
        )

    variedades_sheet = _sheet(workbook, SHEET_VARIEDADES)
    if variedades_sheet is not None:
        variedades = _read_variedades(variedades_sheet)
    else:
        variedades = [Variedad(codigo) for codigo in sorted({o.variedad for o in observaciones})]

    return observaciones, variedades


def escribir_salida(resultado: Resultado) -> bytes:
    workbook = Workbook()

    relativos = workbook.active
    relativos.title = "Relativos"
    meses = resultado.meses[1:]
    relativos.append(["VARIEDAD", *meses])
    by_variedad: dict[str, dict[str, float | None]] = {}
    for r in resultado.relativos:
        by_variedad.setdefault(r.variedad, {})[r.mes] = r.valor
    for variedad in sorted(by_variedad):
        relativos.append([variedad, *(by_variedad[variedad].get(m) for m in meses)])

    precios = workbook.create_sheet("Precios")
    precios.append(["VARIEDAD", "INFORMANTE", "TIPO", "MES", "OBSERVADO", "IMPUTADO", "UTILIZADO"])
    keys = sorted(set(resultado.observados) | set(resultado.imputados) | set(resultado.utilizados))
    for variedad, informante in keys:
        for mes in resultado.meses:
            values = [
                resultado.observados.get((variedad, informante), {}).get(mes),
                resultado.imputados.get((variedad, informante), {}).get(mes),
                resultado.utilizados.get((variedad, informante), {}).get(mes),
            ]
            if any(v is not None for v in values):
                precios.append([variedad, informante, resultado.tipos[informante], mes, *values])

    imputaciones = workbook.create_sheet("Imputaciones")
    imputaciones.append(["VARIEDAD", "INFORMANTE", "MES", "PRECIO", "NIVEL", "GRUPO", "FACTOR"])
    for i in resultado.imputaciones:
        imputaciones.append([i.variedad, i.informante, i.mes, i.precio, i.nivel, i.grupo, i.factor])

    sin_imputar = workbook.create_sheet("Sin imputar")
    sin_imputar.append(["VARIEDAD", "INFORMANTE", "MES"])
    for s in resultado.sin_imputar:
        sin_imputar.append([s.variedad, s.informante, s.mes])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
