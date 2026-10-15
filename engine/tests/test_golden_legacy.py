"""Characterization test against the historical series computed by the original scripts.

The real data is not part of the repository: point SIPC_GOLDEN_DIR to the folder that holds
`Series_precios/SEPTIEMBRE_2020.xlsx`, `Series_variaciones/SEPTIEMBRE_2020.xlsx` and
`Recursos/TRATAMIENTO_ESPECIAL.xlsx`. The test is skipped otherwise.

The original scripts also applied quality changes (Camb.Calidad), which come from the INDEC
export and have no equivalent in this system yet. So only the rows whose whole imputation
chain is free of that artefact are compared; for those, the match must be exact.
"""
import collections
import math
import os
from pathlib import Path

import pytest

from sipc_engine.domain.models import Observacion, Variedad
from sipc_engine.domain.pipeline import calcular

GOLDEN_DIR = os.environ.get("SIPC_GOLDEN_DIR")
pytestmark = pytest.mark.skipif(not GOLDEN_DIR, reason="SIPC_GOLDEN_DIR is not set")

PREFIX_LEN = {1: 8, 2: 6, 3: 5, 4: 4, 5: 3, 6: 1}


def _text(value):
    if value is None:
        return None
    if isinstance(value, (int, float)) and float(value).is_integer():
        return str(int(value))
    return str(value)


def _present(value):
    return value is not None and not (isinstance(value, float) and math.isnan(value))


def _close(a, b, tolerance=1e-6):
    return abs(a - b) <= tolerance * max(1.0, abs(a))


def _sheet(workbook, name):
    rows = list(workbook[name].iter_rows(values_only=True))
    return rows[0], rows[1:]


def _prices(header, rows, id_columns):
    result = {}
    months = [h for h in header[4:]]
    for row in rows:
        key = (_text(row[id_columns[0]]), _text(row[id_columns[1]]))
        result[key] = {m: float(p) for m, p in zip(months, row[4:]) if _present(p)}
    return result


@pytest.fixture(scope="module")
def golden():
    from openpyxl import load_workbook

    root = Path(GOLDEN_DIR)
    series = load_workbook(root / "Series_precios" / "SEPTIEMBRE_2020.xlsx", read_only=True, data_only=True)
    variaciones = load_workbook(root / "Series_variaciones" / "SEPTIEMBRE_2020.xlsx", read_only=True, data_only=True)
    especial = load_workbook(root / "Recursos" / "TRATAMIENTO_ESPECIAL.xlsx", read_only=True, data_only=True)

    special = {
        _text(row[0])
        for sheet in especial.worksheets
        for row in sheet.iter_rows(values_only=True)
        if row and row[0] is not None
    }

    header, rows = _sheet(series, "Precios")
    months = list(header[4:])
    code12 = {_text(r[0]): r[1] for r in rows}
    observaciones = [
        Observacion(_text(r[0]), _text(r[2]), r[3], m, float(p))
        for r in rows
        for m, p in zip(months, r[4:])
        if _present(p)
    ]

    def chain(code):
        return (code, code[:6], code[:5], code[:4], code[:3], code[:1]) if code else ()

    variedades = [Variedad(v, chain(c), v in special) for v, c in code12.items()]
    resultado = calcular(observaciones, variedades, months[0], months[-1])

    h, r = _sheet(series, "Precios imputados")
    legacy_imputados = _prices(h, r, (0, 2))
    h, r = _sheet(series, "Precios utilizados")
    legacy_utilizados = _prices(h, r, (0, 1))

    vh, vrows = _sheet(variaciones, "Variaciones")
    legacy_relativos = {}
    for row in vrows:
        for column, value in zip(vh, row):
            if isinstance(column, str) and column.startswith("20") and _present(value):
                legacy_relativos[(_text(row[0]), column)] = float(value)

    observed = collections.defaultdict(dict)
    for o in observaciones:
        observed[(o.variedad, o.informante)][o.mes] = o.precio

    # Rows where the original scripts changed an observed price (quality change) are tainted,
    # and so is every group that contains one.
    tainted_rows = {
        key
        for key, by_month in observed.items()
        for m, price in by_month.items()
        if legacy_imputados.get(key, {}).get(m) is None or not _close(legacy_imputados[key][m], price)
    }
    tainted_codes = {code12[k[0]] for k in tainted_rows}
    tainted_by_len = {n: {c[:n] for c in tainted_codes if c} for n in set(PREFIX_LEN.values())}
    tainted_variedades = {k[0] for k in tainted_rows}

    levels = collections.defaultdict(set)
    for imputacion in resultado.imputaciones:
        levels[(imputacion.variedad, imputacion.informante)].add(imputacion.nivel)

    def is_clean(key):
        code = code12[key[0]]
        if not code or key[0] in special or key in tainted_rows:
            return False
        for level in levels.get(key, ()):
            if level == 0:
                if key[0] in tainted_variedades:
                    return False
            elif level in PREFIX_LEN:
                if code[: PREFIX_LEN[level]] in tainted_by_len[PREFIX_LEN[level]]:
                    return False
            else:
                return False
        return True

    return {
        "resultado": resultado,
        "months": months,
        "special": special,
        "imputados": legacy_imputados,
        "utilizados": legacy_utilizados,
        "relativos": legacy_relativos,
        "is_clean": is_clean,
        "tainted_variedades": tainted_variedades,
    }


def _compare(legacy, engine, months, keys, is_clean):
    compared = mismatches = 0
    for key in keys:
        if not is_clean(key):
            continue
        for m in months:
            a, b = legacy.get(key, {}).get(m), engine.get(key, {}).get(m)
            if a is None and b is None:
                continue
            compared += 1
            if a is None or b is None or not _close(a, b):
                mismatches += 1
    return compared, mismatches


def test_imputed_prices_match_the_original_scripts(golden):
    keys = set(golden["imputados"]) | set(golden["resultado"].imputados)

    compared, mismatches = _compare(
        golden["imputados"], golden["resultado"].imputados, golden["months"], keys, golden["is_clean"]
    )

    assert compared > 10_000
    assert mismatches == 0


def test_used_prices_match_the_original_scripts(golden):
    keys = set(golden["utilizados"]) | set(golden["resultado"].utilizados)
    clean_variedades = {
        k[0]
        for k in set(golden["imputados"]) | set(golden["resultado"].imputados)
        if golden["is_clean"](k)
    }

    compared, mismatches = _compare(
        golden["utilizados"],
        golden["resultado"].utilizados,
        golden["months"],
        {k for k in keys if k[0] in clean_variedades},
        golden["is_clean"],
    )

    assert compared > 5_000
    assert mismatches == 0


def test_relatives_match_the_original_scripts(golden):
    all_keys = set(golden["imputados"]) | set(golden["resultado"].imputados) | set(golden["utilizados"])
    by_variedad = collections.defaultdict(list)
    for key in all_keys:
        by_variedad[key[0]].append(key)
    clean = {
        v
        for v, keys in by_variedad.items()
        if v not in golden["special"] and all(golden["is_clean"](k) for k in keys)
    }

    engine = {(r.variedad, r.mes): r.valor for r in golden["resultado"].relativos if r.valor is not None}
    legacy = golden["relativos"]
    cells = [k for k in set(legacy) | set(engine) if k[0] in clean]

    assert len(cells) > 500
    for key in cells:
        assert key in legacy and key in engine, key
        assert _close(legacy[key], engine[key], 1e-9), key
