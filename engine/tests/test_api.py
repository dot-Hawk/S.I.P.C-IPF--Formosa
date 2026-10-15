import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from sipc_engine.api.app import create_app

KEY = "test-key"
HEADERS = {"X-API-Key": KEY}


@pytest.fixture()
def client():
    return TestClient(create_app(api_key=KEY))


def payload(**overrides):
    body = {
        "mesBase": "2026-01",
        "mesObjetivo": "2026-02",
        "variedades": [{"codigo": "V1"}],
        "observaciones": [
            {"variedad": "V1", "informante": "i1", "tipo": "TRAD", "mes": "2026-01", "precio": 100},
            {"variedad": "V1", "informante": "i2", "tipo": "TRAD", "mes": "2026-01", "precio": 200},
            {"variedad": "V1", "informante": "i1", "tipo": "TRAD", "mes": "2026-02", "precio": 110},
        ],
    }
    body.update(overrides)
    return body


def workbook_bytes(precios, variedades=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Precios"
    ws.append(["VARIEDAD", "INFORMANTE", "TIPO", "MES", "PRECIO"])
    for row in precios:
        ws.append(row)
    if variedades is not None:
        vs = wb.create_sheet("Variedades")
        vs.append(["VARIEDAD", "VARIEDAD_12", "TRATAMIENTO_ESPECIAL"])
        for row in variedades:
            vs.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


EXCEL_PRECIOS = [
    ["V1", "i1", "TRAD", "2026-01", 100],
    ["V1", "i2", "TRAD", "2026-01", 200],
    ["V1", "i1", "TRAD", "2026-02", 110],
]


class TestAuth:
    def test_health_is_public(self, client):
        response = client.get("/health")

        assert response.status_code == 200
        assert response.json()["status"] == "ok"

    def test_calculo_requires_the_api_key(self, client):
        assert client.post("/v1/calculo", json=payload()).status_code == 401
        assert client.post("/v1/calculo", json=payload(), headers={"X-API-Key": "nope"}).status_code == 401

    def test_an_engine_without_a_configured_key_refuses_to_calculate(self):
        client = TestClient(create_app(api_key=None))

        response = client.post("/v1/calculo", json=payload(), headers=HEADERS)

        assert response.status_code == 503


class TestCalculoJson:
    def test_returns_relatives_prices_and_imputations(self, client):
        response = client.post("/v1/calculo", json=payload(), headers=HEADERS)

        assert response.status_code == 200
        body = response.json()
        assert body["meses"] == ["2026-01", "2026-02"]
        assert body["relativos"] == [
            {"variedad": "V1", "mes": "2026-02", "valor": pytest.approx(1.1), "nSuper": 0, "nTrad": 2}
        ]
        imputado = next(p for p in body["precios"] if p["informante"] == "i2" and p["mes"] == "2026-02")
        assert imputado["observado"] is None
        assert imputado["imputado"] == pytest.approx(220)
        assert body["imputaciones"][0]["nivel"] == 0
        assert body["sinImputar"] == []
        assert body["resumen"]["observaciones"] == 3

    def test_accepts_parameters(self, client):
        body = payload(
            observaciones=[
                {"variedad": "V1", "informante": "s", "tipo": "SUPER", "mes": "2026-01", "precio": 100},
                {"variedad": "V1", "informante": "t", "tipo": "TRAD", "mes": "2026-01", "precio": 100},
                {"variedad": "V1", "informante": "s", "tipo": "SUPER", "mes": "2026-02", "precio": 120},
                {"variedad": "V1", "informante": "t", "tipo": "TRAD", "mes": "2026-02", "precio": 110},
            ],
            parametros={"pesoSuper": 0.5, "pesoTrad": 0.5},
        )

        response = client.post("/v1/calculo", json=body, headers=HEADERS)

        assert response.json()["relativos"][0]["valor"] == pytest.approx((1.2 * 1.1) ** 0.5)

    def test_domain_errors_become_422_with_the_reason(self, client):
        body = payload(mesObjetivo="2026-03")

        response = client.post("/v1/calculo", json=body, headers=HEADERS)

        assert response.status_code == 422
        assert "2026-03" in response.json()["detail"]

    def test_malformed_input_is_rejected(self, client):
        response = client.post("/v1/calculo", json={"mesBase": "2026-01"}, headers=HEADERS)

        assert response.status_code == 422


class TestCalculoExcel:
    def test_reads_the_long_format_sheet(self, client):
        response = client.post(
            "/v1/calculo/excel",
            files={"file": ("precios.xlsx", workbook_bytes(EXCEL_PRECIOS))},
            data={"mesBase": "2026-01"},
            headers=HEADERS,
        )

        assert response.status_code == 200
        body = response.json()
        assert body["meses"] == ["2026-01", "2026-02"]  # target defaults to the last month in the file
        assert body["relativos"][0]["valor"] == pytest.approx(1.1)

    def test_builds_the_imputation_hierarchy_and_special_flag_from_the_variedades_sheet(self, client):
        precios = EXCEL_PRECIOS + [
            ["E1", "i1", "TRAD", "2026-01", 5],
            ["E1", "i1", "TRAD", "2026-02", 9],
        ]
        variedades = [["V1", "A0111111", None], ["E1", "A0111112", "SI"]]

        response = client.post(
            "/v1/calculo/excel",
            files={"file": ("precios.xlsx", workbook_bytes(precios, variedades))},
            data={"mesBase": "2026-01", "mesObjetivo": "2026-02"},
            headers=HEADERS,
        )

        body = response.json()
        assert body["especiales"] == ["E1"]
        assert all(r["variedad"] != "E1" for r in body["relativos"])

    def test_can_return_an_excel_workbook(self, client):
        response = client.post(
            "/v1/calculo/excel?formato=xlsx",
            files={"file": ("precios.xlsx", workbook_bytes(EXCEL_PRECIOS))},
            data={"mesBase": "2026-01"},
            headers=HEADERS,
        )

        assert response.status_code == 200
        assert "spreadsheetml" in response.headers["content-type"]
        workbook = load_workbook(io.BytesIO(response.content))
        assert {"Relativos", "Precios", "Imputaciones", "Sin imputar"} <= set(workbook.sheetnames)
        relativos = list(workbook["Relativos"].iter_rows(values_only=True))
        assert relativos[0] == ("VARIEDAD", "2026-02")
        assert relativos[1][0] == "V1"
        assert relativos[1][1] == pytest.approx(1.1)

    def test_rejects_a_file_without_the_precios_sheet(self, client):
        wb = Workbook()
        wb.active.title = "Otra"
        buffer = io.BytesIO()
        wb.save(buffer)

        response = client.post(
            "/v1/calculo/excel",
            files={"file": ("x.xlsx", buffer.getvalue())},
            data={"mesBase": "2026-01"},
            headers=HEADERS,
        )

        assert response.status_code == 422
        assert "Precios" in response.json()["detail"]

    def test_rejects_a_file_that_is_not_an_excel(self, client):
        response = client.post(
            "/v1/calculo/excel",
            files={"file": ("x.xlsx", b"not an excel")},
            data={"mesBase": "2026-01"},
            headers=HEADERS,
        )

        assert response.status_code == 422

    def test_reports_the_row_of_an_invalid_value(self, client):
        precios = [["V1", "i1", "TRAD", "2026-01", "abc"]]

        response = client.post(
            "/v1/calculo/excel",
            files={"file": ("x.xlsx", workbook_bytes(precios))},
            data={"mesBase": "2026-01"},
            headers=HEADERS,
        )

        assert response.status_code == 422
        assert "fila 2" in response.json()["detail"]
